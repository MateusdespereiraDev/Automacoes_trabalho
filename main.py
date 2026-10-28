import pyautogui as tempo 
import pyautogui as posicao
from openpyxl import load_workbook
import pyperclip
#bibliotecas usadas

#setup para o python ler o documento 
ARQUIVO = "C:\\Users\\ML Placas\\Downloads\\tags_mantidas.xlsx"
arquivo = load_workbook(ARQUIVO)
dados = arquivo.active

#detalhes de controle do código
i=1
posicao.PAUSE = 0.5

"""botao_adcionar_lead = posicao.position(x=1529, y=318)"""
click_adicionar_tag = posicao.position(x=1291, y=184)
click_criar = posicao.position(x=803, y=547)


tempo.sleep(2)
#lê linha por linha e pega seu valor
for i in dados.iter_rows(min_row=49, values_only=True):
        tag = i[0]
        #daqui para baixo vai ser feita a importação, é uma série de teclas e ações
        posicao.click(click_adicionar_tag)
        pyperclip.copy(str(tag))
        posicao.hotkey("ctrl", "v")
        posicao.press("tab", presses=2)
        
   
 
        """posicao.press("tab", presses=3)
        pyperclip.copy("impo")
        posicao.hotkey("ctrl", "v")"""
       
      

        """posicao.click(click_adicionar_tag)"""
        posicao.click(click_criar)
        tempo.sleep(0.5)

